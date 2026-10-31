"""把 S4 回测（默认：50 万、半仓、交易所保证金 12%、手续费交易所加一分钱、复利）导出为 Excel 报告与资金曲线图。

用法：python export_s4_excel.py --csv data/JM_weighted_user.csv
输出：results/S4回测报告_50万半仓.xlsx、results/S4资金曲线_50万半仓.png
Excel 中“核心数据”“年度收益”均为公式，引用“逐笔交易”“资金曲线”两张表（这两张表是回测程序输出的数值）。
"""
import argparse
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from jm_strategy.backtest import Account, metrics, run_backtest
from jm_strategy.data import load_csv
from jm_strategy.library import backtest_params, base_series, build_signals

S4 = dict(sar="s1", filt="stm4", exit="sar:slow", stop="atr2.5")
BASE = dict(capital=500_000, sizing="margin", pos_frac=0.5, margin_rate=0.12,
            fee_rate=1e-4, fee_per_lot=0.01, fee_rate_today=2e-4)
REASONS = {"指标离场": "慢SAR翻转离场", "止损": "止损", "回测结束": "回测结束(仍持仓)"}

FONT = "Arial"
BLUE, GREEN = "0000FF", "008000"
HEAD_FILL = PatternFill("solid", fgColor="E8E8E8")
THIN = Side(style="thin", color="C8C8C8")


def font(bold=False, color="000000", size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)


def header(ws, row, titles, widths=None):
    for j, t in enumerate(titles, start=1):
        c = ws.cell(row=row, column=j, value=t)
        c.font, c.fill = font(bold=True), HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=THIN)
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w


def run(df, sig, **kw):
    acct = Account(**{**BASE, **kw})
    eq, tr, _ = run_backtest(df, backtest_params(sig), acct, sig=sig,
                             **{k: v for k, v in kw.items() if k in ("start", "end")})
    return eq, tr, acct


def scenario_rows(df, sig):
    """其他口径对照（Python 回测输出的数值）。"""
    cases = [("样本内 2013–2020（同设置）", {}, dict(end="2020-12-31")),
             ("样本外 2021–2026（同设置，从 50 万重新开始）", {}, dict(start="2021-01-01")),
             ("不复利（每次按 50 万计算手数）", dict(compound=False), {}),
             ("保证金 15%", dict(margin_rate=0.15), {}),
             ("保证金 20%", dict(margin_rate=0.20), {}),
             ("按风险定仓：每笔止损亏损 = 权益 5%", dict(sizing="risk", risk_per_trade=0.05, max_leverage=5), {}),
             ("按风险定仓：每笔止损亏损 = 权益 8%", dict(sizing="risk", risk_per_trade=0.08, max_leverage=5), {})]
    rows = []
    for name, kw, rng in cases:
        acct = Account(**{**BASE, **kw})
        eq, tr, _ = run_backtest(df, backtest_params(sig), acct, sig=sig, **rng)
        m = metrics(eq, tr, acct.capital)
        prev = eq.shift(1).fillna(acct.capital)
        pct = pd.Series([t.pnl / prev[t.entry_date] for t in tr.itertuples()])
        rows.append([name, m["年化收益率"], m["最大回撤"], m["交易次数"], m["胜率"], m["盈亏比(均盈/均亏)"],
                     pct[pct > 0].mean() / -pct[pct <= 0].mean(), pct.min(), eq.iloc[-1]])
    return rows


def build_workbook(df, eq, tr, scen, path):
    wb = Workbook()

    # ---------------- 参数说明 ----------------
    ws = wb.active
    ws.title = "参数说明"
    ws["A1"] = "S4 焦煤加权回测：参数与交易规则"
    ws["A1"].font = font(bold=True, size=13)
    header(ws, 3, ["参数", "数值", "说明 / 来源"], [30, 16, 80])
    params = [
        ("初始资金（元）", 500000, "#,##0", "用户设定"),
        ("每次开仓保证金占权益比例", 0.5, "0%", "用户设定（半仓），每次按开仓时的账户权益计算手数（复利）"),
        ("保证金比例", 0.12, "0%", "交易所标准，期货公司不上浮；12% 来自网络搜索结果，未能在大商所官网核实"),
        ("合约乘数（吨/手）", 60, "0", "大商所焦煤合约"),
        ("手续费率（每边，按成交额）", 0.0001, "0.00%", "交易所标准万分之一（网络资料，未核实）"),
        ("平今平仓手续费率", 0.0002, "0.00%", "当天开当天平；2025-08 起可能为万分之二（网络资料，未核实）"),
        ("每边每手另加手续费（元）", 0.01, "0.00", "用户设定：交易所加一分钱"),
        ("滑点（每边，元/吨）", 0.5, "0.0", "按 1 个最小变动价位假设"),
        ("回测数据", "焦煤加权日线", "@", "用户提供，2013-03-22（上市首日）至 2026-10-29，共 3304 根 K 线"),
    ]
    for i, (k, v, fmt, note) in enumerate(params, start=4):
        ws.cell(row=i, column=1, value=k).font = font()
        c = ws.cell(row=i, column=2, value=v)
        c.font, c.number_format = font(color=BLUE), fmt
        ws.cell(row=i, column=3, value=note).font = font()
    r = 4 + len(params) + 1
    ws.cell(row=r, column=1, value="注意").font = font(bold=True)
    ws.cell(row=r, column=3, value="“逐笔交易”“资金曲线”两张表是 Python 回测程序（run_s4.py）的输出数值；"
            "修改上面的参数不会重新回测，需要重跑程序。“核心数据”“年度收益”由公式从这两张表计算。").font = font()
    ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True)
    r += 2
    ws.cell(row=r, column=1, value="交易规则（多空对称，以做多为例）").font = font(bold=True, size=11)
    rules = [
        ("用到的指标", "灵敏 SAR(5, 3%, 30%) 抓转折；指标“1”（简化 STM = (MA5−MA20) + (DIF−DEA)×5 + (RSI14−50)/2）确认方向；"
                   "慢 SAR(10, 0.5%, 5%) 跟踪止盈；ATR(14) 定止损"),
        ("准备", "灵敏 SAR 还在 K 线上方，指标“1”已在零轴上方或绿柱快速缩短，价格逼近 SAR 点"),
        ("进场", "同一根日 K 线收盘时：灵敏 SAR 由空翻多，且指标“1” > 0 → 次日开盘买入。翻转当天指标“1”仍 < 0 则放弃这次信号"),
        ("做空", "灵敏 SAR 由多翻空，且指标“1” < 0 → 次日开盘卖出"),
        ("止损", "入场价 − 2.5 × ATR(14)（取信号当天 ATR），固定不动；盘中触及即平，跳空越过按开盘价成交"),
        ("止盈/离场", "持仓期间慢 SAR 由多翻空（入场后发生的翻转）→ 次日开盘平仓；没有固定目标价"),
        ("持仓期间", "不加仓、不反手，只看慢 SAR 与止损；平仓后空仓等待下一个信号"),
        ("仓位", "手数 = 开仓时权益 × 50% ÷（开仓价 × 60 × 保证金比例），向下取整"),
    ]
    for k, v in rules:
        r += 1
        ws.cell(row=r, column=1, value=k).font = font(bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = ws.cell(row=r, column=2, value=v)
        c.font, c.alignment = font(), Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30

    # ---------------- 逐笔交易 ----------------
    wt = wb.create_sheet("逐笔交易")
    header(wt, 1, ["序号", "方向", "开仓日期", "开仓价", "平仓日期", "平仓价", "手数", "持仓天数",
                   "净盈亏（元）", "开仓前权益（元）", "盈亏占权益", "R 倍数", "离场原因", "名义杠杆",
                   "连亏计数", "连赢计数"],
           [6, 6, 12, 10, 12, 10, 8, 9, 15, 16, 11, 9, 18, 10, 9, 9])
    wt.cell(row=1, column=9).comment = Comment("含开仓、平仓手续费和滑点", "Claude")
    wt.cell(row=1, column=12).comment = Comment("(平仓价 − 开仓价) × 方向 ÷ 初始止损距离（2.5 ATR）", "Claude")
    n = len(tr)
    for i, t in enumerate(tr.itertuples(), start=2):
        vals = [i - 1, t.side, t.entry_date.to_pydatetime(), t.entry, t.exit_date.to_pydatetime(), t.exit,
                t.lots, t.bars, t.pnl]
        fmts = ["0", "@", "yyyy-mm-dd", "0.0", "yyyy-mm-dd", "0.0", "#,##0", "0", "#,##0;-#,##0"]
        for j, (v, f) in enumerate(zip(vals, fmts), start=1):
            c = wt.cell(row=i, column=j, value=v)
            c.font, c.number_format = font(), f
        wt.cell(row=i, column=10, value="=参数说明!$B$4" if i == 2 else f"=J{i - 1}+I{i - 1}")
        wt.cell(row=i, column=10).font = font(color=GREEN if i == 2 else "000000")
        wt.cell(row=i, column=11, value=f"=I{i}/J{i}")
        wt.cell(row=i, column=12, value=round(t.R, 3))
        wt.cell(row=i, column=13, value=REASONS.get(t.reason, t.reason))
        wt.cell(row=i, column=14, value=f"=G{i}*D{i}*参数说明!$B$7/J{i}")
        wt.cell(row=i, column=15, value=f"=IF(I{i}<=0,1,0)" if i == 2 else f"=IF(I{i}<=0,O{i - 1}+1,0)")
        wt.cell(row=i, column=16, value=f"=IF(I{i}>0,1,0)" if i == 2 else f"=IF(I{i}>0,P{i - 1}+1,0)")
        for j, f in ((10, "#,##0"), (11, "0.0%;-0.0%"), (12, "0.00"), (14, '0.00"倍"'), (15, "0"), (16, "0")):
            wt.cell(row=i, column=j).number_format = f
        for j in (11, 12, 13, 14, 15, 16):
            wt.cell(row=i, column=j).font = font()
    wt.freeze_panes = "C2"
    last = n + 1

    # ---------------- 资金曲线 ----------------
    we = wb.create_sheet("资金曲线")
    header(we, 1, ["日期", "焦煤加权收盘", "账户权益（元）", "历史最高权益（元）", "回撤", "年份"], [12, 13, 18, 18, 10, 8])
    we.cell(row=1, column=3).comment = Comment("每日收盘按持仓浮动盈亏计算（回测程序输出）", "Claude")
    for i, (d, e) in enumerate(eq.items(), start=2):
        we.cell(row=i, column=1, value=d.to_pydatetime()).number_format = "yyyy-mm-dd"
        we.cell(row=i, column=2, value=float(df.close[d])).number_format = "0.0"
        we.cell(row=i, column=3, value=round(float(e), 2)).number_format = "#,##0"
        we.cell(row=i, column=4, value="=C2" if i == 2 else f"=MAX(D{i - 1},C{i})").number_format = "#,##0"
        we.cell(row=i, column=5, value=f"=C{i}/D{i}-1").number_format = "0.0%"
        we.cell(row=i, column=6, value=f"=YEAR(A{i})").number_format = "0"
    we.freeze_panes = "A2"
    elast = len(eq) + 1

    # ---------------- 核心数据 ----------------
    wc = wb.create_sheet("核心数据", 0)
    wc["A1"] = "S4 回测核心数据（50 万本金，半仓，保证金 12%，手续费交易所加一分钱，复利）"
    wc["A1"].font = font(bold=True, size=13)
    header(wc, 3, ["指标", "数值", "说明"], [30, 20, 70])
    T, E = "逐笔交易", "资金曲线"
    I, K, H, B, M, N_ = (f"{T}!$I$2:$I${last}", f"{T}!$K$2:$K${last}", f"{T}!$H$2:$H${last}",
                         f"{T}!$B$2:$B${last}", f"{T}!$M$2:$M${last}", f"{T}!$N$2:$N${last}")
    rows = [
        ("初始资金（元）", "=参数说明!B4", "#,##0", ""),
        ("期末权益（元）", f"={E}!C{elast}", "#,##0", "2026-10-29 收盘，含当时持仓的浮动盈亏"),
        ("总收益率", "=B5/B4-1", "0.0%", ""),
        ("回测年数", f"=({E}!A{elast}-{E}!A2)/365.25", "0.00", "2013-03-22 至 2026-10-29"),
        ("年化收益率", "=(B5/B4)^(1/B7)-1", "0.0%", ""),
        ("最大回撤", f"=MIN({E}!E2:E{elast})", "0.0%", "按每日收盘权益计算"),
        ("最大回撤谷底日期", f"=INDEX({E}!A2:A{elast},MATCH(B9,{E}!E2:E{elast},0))", "yyyy-mm-dd", ""),
        ("卡玛比率（年化 ÷ 最大回撤）", "=B8/ABS(B9)", "0.00", ""),
        ("交易次数", f"=COUNT({I})", "0", ""),
        ("年均交易次数", "=B12/B7", "0.0", ""),
        ("盈利笔数", f'=COUNTIF({I},">0")', "0", ""),
        ("亏损笔数", f'=COUNTIF({I},"<=0")', "0", ""),
        ("胜率", "=B14/B12", "0.0%", ""),
        ("平均盈利（元）", f'=AVERAGEIF({I},">0")', "#,##0", ""),
        ("平均亏损（元）", f'=AVERAGEIF({I},"<=0")', "#,##0;-#,##0", ""),
        ("盈亏比（按金额）", "=B17/ABS(B18)", "0.00", "复利时后期仓位巨大，金额口径被最后几年主导，参考价值有限"),
        ("平均盈利占当时权益", f'=AVERAGEIF({I},">0",{K})', "0.0%", ""),
        ("平均亏损占当时权益", f'=AVERAGEIF({I},"<=0",{K})', "0.0%;-0.0%", ""),
        ("盈亏比（按占权益比例）", "=B20/ABS(B21)", "0.00", "更能反映策略本身"),
        ("利润因子（总盈利 ÷ 总亏损）", f'=SUMIF({I},">0")/ABS(SUMIF({I},"<=0"))', "0.00", ""),
        ("每笔平均盈亏（元）", f"=AVERAGE({I})", "#,##0", ""),
        ("单笔最大盈利（元）", f"=MAX({I})", "#,##0", ""),
        ("单笔最大盈利占权益", f"=MAX({K})", "0.0%", ""),
        ("单笔最大亏损（元）", f"=MIN({I})", "#,##0;-#,##0", ""),
        ("单笔最大亏损占权益", f"=MIN({K})", "0.0%;-0.0%", ""),
        ("最大连续亏损（笔）", f"=MAX({T}!$O$2:$O${last})", "0", ""),
        ("最大连续盈利（笔）", f"=MAX({T}!$P$2:$P${last})", "0", ""),
        ("平均持仓天数（交易日）", f"=AVERAGE({H})", "0.0", ""),
        ("盈利单平均持仓天数", f'=AVERAGEIF({I},">0",{H})', "0.0", ""),
        ("亏损单平均持仓天数", f'=AVERAGEIF({I},"<=0",{H})', "0.0", ""),
        ("慢SAR翻转离场笔数", f'=COUNTIF({M},"慢SAR翻转离场")', "0", ""),
        ("止损离场笔数", f'=COUNTIF({M},"止损")', "0", ""),
        ("多单笔数", f'=COUNTIF({B},"多")', "0", ""),
        ("多单胜率", f'=COUNTIFS({B},"多",{I},">0")/B36', "0.0%", ""),
        ("多单盈亏合计（元）", f'=SUMIF({B},"多",{I})', "#,##0;-#,##0", ""),
        ("空单笔数", f'=COUNTIF({B},"空")', "0", ""),
        ("空单胜率", f'=COUNTIFS({B},"空",{I},">0")/B39', "0.0%", ""),
        ("空单盈亏合计（元）", f'=SUMIF({B},"空",{I})', "#,##0;-#,##0", ""),
        ("平均名义杠杆（合约市值 ÷ 权益）", f"=AVERAGE({N_})", '0.00"倍"', ""),
        ("最高名义杠杆", f"=MAX({N_})", '0.00"倍"', ""),
        ("对账：逐笔盈亏合计 − 权益增加", f"=SUM({I})-(B5-B4)", "#,##0.00", "应为 0（期末无持仓时）"),
    ]
    for i, (k, f, fmt, note) in enumerate(rows, start=4):
        wc.cell(row=i, column=1, value=k).font = font()
        c = wc.cell(row=i, column=2, value=f)
        c.font, c.number_format = font(color=GREEN if "!" in f else "000000"), fmt
        wc.cell(row=i, column=3, value=note).font = font()
    r = 4 + len(rows) + 1
    wc.cell(row=r, column=1, value="其他口径对照（Python 回测输出的数值）").font = font(bold=True, size=11)
    r += 1
    cols = ["口径", "年化收益率", "最大回撤", "交易次数", "胜率", "盈亏比（金额）", "盈亏比（占权益）", "单笔最大亏损", "期末权益（元）"]
    for j, t in enumerate(cols, start=1):
        c = wc.cell(row=r, column=j, value=t)
        c.font, c.fill = font(bold=True), HEAD_FILL
    for col, w in zip("DEFGHI", [9, 9, 13, 14, 13, 16]):
        wc.column_dimensions[col].width = w
    fmts = ["@", "0.0%", "0.0%;-0.0%", "0", "0.0%", "0.00", "0.00", "0.0%;-0.0%", "#,##0"]
    for row in scen:
        r += 1
        for j, (v, f) in enumerate(zip(row, fmts), start=1):
            c = wc.cell(row=r, column=j, value=v)
            c.font, c.number_format = font(), f

    # ---------------- 年度收益 ----------------
    wy = wb.create_sheet("年度收益", 1)
    header(wy, 1, ["年份", "年末日期", "年初权益（元）", "年末权益（元）", "年收益率", "开仓笔数", "当年平仓盈亏（元）"],
           [8, 12, 18, 18, 11, 10, 18])
    years = sorted(eq.index.year.unique())
    for i, y in enumerate(years, start=2):
        wy.cell(row=i, column=1, value=int(y)).number_format = "0"
        wy.cell(row=i, column=2, value=f"=_xlfn.MAXIFS({E}!$A$2:$A${elast},{E}!$F$2:$F${elast},A{i})").number_format = "yyyy-mm-dd"
        wy.cell(row=i, column=3, value="=参数说明!$B$4" if i == 2 else f"=D{i - 1}").number_format = "#,##0"
        wy.cell(row=i, column=4, value=f"=INDEX({E}!$C$2:$C${elast},MATCH(B{i},{E}!$A$2:$A${elast},0))").number_format = "#,##0"
        wy.cell(row=i, column=5, value=f"=D{i}/C{i}-1").number_format = "0.0%;-0.0%"
        wy.cell(row=i, column=6, value=f'=COUNTIFS({T}!$C$2:$C${last},">="&DATE(A{i},1,1),{T}!$C$2:$C${last},"<="&DATE(A{i},12,31))').number_format = "0"
        wy.cell(row=i, column=7, value=f'=SUMIFS({T}!$I$2:$I${last},{T}!$E$2:$E${last},">="&DATE(A{i},1,1),{T}!$E$2:$E${last},"<="&DATE(A{i},12,31))').number_format = "#,##0;-#,##0"
        for j in range(1, 8):
            wy.cell(row=i, column=j).font = font(color=GREEN if j in (2, 4, 6, 7) else "000000")
    wy.freeze_panes = "A2"

    # ---------------- 图表 ----------------
    wg = wb.create_sheet("图表", 2)
    ch = LineChart()
    ch.title = "账户权益（元，对数坐标）"
    ch.y_axis.scaling.logBase = 10
    ch.y_axis.title = "元"
    ch.x_axis.number_format = "yyyy"
    ch.height, ch.width = 9, 22
    ch.legend = None
    ch.add_data(Reference(we, min_col=3, min_row=1, max_row=elast), titles_from_data=True)
    ch.set_categories(Reference(we, min_col=1, min_row=2, max_row=elast))
    ch.series[0].graphicalProperties.line.solidFill = "2A78D6"
    ch.series[0].graphicalProperties.line.width = 15000
    ch.series[0].smooth = False
    wg.add_chart(ch, "A1")
    dd = LineChart()
    dd.title = "回撤"
    dd.y_axis.number_format = "0%"
    dd.x_axis.number_format = "yyyy"
    dd.height, dd.width = 7, 22
    dd.legend = None
    dd.add_data(Reference(we, min_col=5, min_row=1, max_row=elast), titles_from_data=True)
    dd.set_categories(Reference(we, min_col=1, min_row=2, max_row=elast))
    dd.series[0].graphicalProperties.line.solidFill = "E34948"
    dd.series[0].graphicalProperties.line.width = 12000
    wg.add_chart(dd, "A20")

    wb.save(path)


def plot(eq, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter
    plt.rcParams["font.family"] = ["WenQuanYi Zen Hei", "DejaVu Sans"]
    surface, ink, ink2, grid = "#fcfcfb", "#0b0b0b", "#52514e", "#e4e3df"
    fig, ax = plt.subplots(2, 1, figsize=(11, 7), sharex=True, facecolor=surface,
                           gridspec_kw=dict(height_ratios=[3, 1.2]))
    for a in ax:
        a.set_facecolor(surface)
        for s in ("top", "right"):
            a.spines[s].set_visible(False)
        for s in ("left", "bottom"):
            a.spines[s].set_color(grid)
        a.tick_params(colors=ink2, labelsize=9)
        a.grid(color=grid, linewidth=0.8)
        a.set_axisbelow(True)
    ax[0].plot(eq.index, eq / 1e4, color="#2a78d6", linewidth=2)
    ax[0].set_yscale("log")
    ax[0].yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:,.0f} 万"))
    ax[0].annotate(f"{eq.iloc[-1] / 1e8:.2f} 亿", (eq.index[-1], eq.iloc[-1] / 1e4), xytext=(6, 0),
                   textcoords="offset points", va="center", fontsize=9, color=ink)
    ax[0].set_title("S4 账户权益（50 万起，半仓，保证金 12%，复利；对数坐标）", color=ink, fontsize=12, loc="left")
    d = eq / eq.cummax() - 1
    ax[1].fill_between(d.index, d, 0, color="#e34948", alpha=0.35, linewidth=0)
    ax[1].plot(d.index, d, color="#e34948", linewidth=1)
    ax[1].yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:.0%}"))
    ax[1].set_title("回撤", color=ink, fontsize=11, loc="left")
    i = d.idxmin()
    ax[1].annotate(f"最大回撤 {d.min():.1%}（{i.date()}）", (i, d.min()), xytext=(8, 4), textcoords="offset points",
                   fontsize=9, color=ink)
    fig.tight_layout()
    fig.savefig(path, dpi=120, facecolor=surface)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default="data/JM_weighted_user.csv")
    ap.add_argument("--out", default="results")
    a = ap.parse_args()
    df = load_csv(a.csv)
    sig = build_signals(base_series(df), "S4", S4)
    eq, tr, _ = run_backtest(df, backtest_params(sig), Account(**BASE), sig=sig)
    os.makedirs(a.out, exist_ok=True)
    xlsx = os.path.join(a.out, "S4回测报告_50万半仓.xlsx")
    build_workbook(df, eq, tr, scenario_rows(df, sig), xlsx)
    plot(eq, os.path.join(a.out, "S4资金曲线_50万半仓.png"))
    print("written", xlsx)


if __name__ == "__main__":
    main()
